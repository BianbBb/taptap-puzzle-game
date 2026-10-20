#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把编剧文档里的 {表情} 标记回填到《剧情节点配置表》的 animationAsset 列。

文档格式（各事件对白文本.docx）：

    陈先生：周老，我就直说了。三十天，整条街清完。您这儿是最后一户。{思考}

* `{...}` 是**说话人表情**（本文档里只有 思考 / 悲伤 / 失落 / 气愤 / 喜悦 / 感动 六种）；
* `【...】` 是舞台提示、`[...]`/`（...）` 里的【好奇】【怀念】是记忆情绪标签，都不算表情；
* 没有 `{}` 的行表示「沿用上一句的表情」，回填成空值，运行时保持当前姿势不切换。

写进表里的是**语义键**（见 EMOTION_TO_ASSET），不是骨骼动画名——
每个角色可用的动画不同（老周没有「喜悦」、陈先生没有「气愤」），
角色 → 动画名的解析放在运行时 BattleMainUI.MapNpcSpinePose 里。

用法：

    # 只出报告，不改表
    python3 GameConfig/Tools/story_import/fill_story_expressions.py --src <docx>

    # 写回《剧情节点配置表.xlsx》
    python3 GameConfig/Tools/story_import/fill_story_expressions.py --src <docx> --apply
"""
import argparse
import collections
import difflib
import os
import re
import sys

import docx
import openpyxl

EMOTION_TO_ASSET = {
    "思考": "think",
    "悲伤": "sad",
    "失落": "sad",
    "气愤": "angry",
    "生气": "angry",
    "喜悦": "happy",
    "开心": "happy",
    "感动": "moved",
    "惊讶": "surprised",
}

SPEAKERS = ("陈先生", "老周", "周老", "陈果", "妻子", "陈父", "AI")
SPEAKER_ALIAS = {"周老": "老周"}

MARK_RE = re.compile(r"[｛{]([^{}｝]{1,6})[｝}]")
# 舞台提示 / 记忆标签 / 括注都不参与文本比对
DROP_RE = re.compile(r"【[^】]*】|\[[^\]]*\]|（[^）]*）|\([^)]*\)")
PUNCT_RE = re.compile(r"[\s，。！？、；：“”‘’\"'…—\-\.\!\?·　]+")

TABLE_NAME = "剧情节点配置表.xlsx"
COL_ID, COL_NPC, COL_TEXT, COL_ANIM = 2, 3, 4, 11   # B / C / D / K
DATA_START_ROW = 5


def normalize(text):
    text = DROP_RE.sub("", text or "")
    text = MARK_RE.sub("", text)
    return PUNCT_RE.sub("", text)


def read_doc(path):
    """把 docx 解析成 [(说话人, 归一化文本, 表情键, 原文)]。"""
    document = docx.Document(path)
    lines = [p.text.strip() for p in document.paragraphs]
    lines = [l for l in lines if l]

    # 有人把一句台词拆成「陈先生：」+ 下一段，先并回去
    merged = []
    for line in lines:
        if merged and re.fullmatch(r"[\u4e00-\u9fa5A-Za-z]{1,8}[：:]", merged[-1]):
            merged[-1] = merged[-1] + line
        else:
            merged.append(line)

    out = []
    for line in merged:
        # 去掉行首的（回复1）（A线）（记忆：初见你）这类标注，它们不是说话人
        bare = re.sub(r"^[（(【\[][^）)】\]]*[）)】\]]\s*", "", line)
        m = re.match(r"^([\u4e00-\u9fa5A-Za-z\$]{1,8})[：:]", bare)
        if not m:
            continue
        speaker = m.group(1)
        if speaker not in SPEAKERS:
            continue
        body = bare[m.end():]
        marks = MARK_RE.findall(body)
        key = None
        if marks:
            emotion = marks[-1]
            key = EMOTION_TO_ASSET.get(emotion)
            if key is None:
                print("  [警告] 未知表情标记 {%s}：%s" % (emotion, line[:40]), file=sys.stderr)
        out.append((SPEAKER_ALIAS.get(speaker, speaker), normalize(body), key, line))
    return out


def load_rows(ws):
    rows = []
    for r in range(DATA_START_ROW, ws.max_row + 1):
        node_id = ws.cell(r, COL_ID).value
        if not node_id:
            continue
        rows.append({
            "row": r,
            "id": str(node_id).strip(),
            "npc": (ws.cell(r, COL_NPC).value or "").strip(),
            "text": (ws.cell(r, COL_TEXT).value or "").strip(),
            "norm": normalize(ws.cell(r, COL_TEXT).value),
            "anim": ws.cell(r, COL_ANIM).value,
        })
    return rows


def align(doc_lines, rows):
    """把文档台词对到表行上：先精确，再「互相包含」，最后模糊。"""
    assigned = {}          # row index -> (key, doc原文)
    used_doc = set()

    by_key = collections.defaultdict(list)
    for i, row in enumerate(rows):
        by_key[(row["npc"], row["norm"])].append(i)

    # 1) 归一化后完全一致
    exact_rows = set()
    for di, (speaker, norm_text, key, raw) in enumerate(doc_lines):
        for i in by_key.get((speaker, norm_text), []):
            if i in assigned:
                continue
            assigned[i] = (key, raw)
            exact_rows.add(i)
            used_doc.add(di)
            break

    # 2) 一句被拆成多行 / 多句并成一行：文本互相包含
    #    （文档一句 → 表里连续多个节点时，这几个节点共用同一表情；
    #      表里一行吞了文档好几句时，后一句的覆盖前一句——表情以句子结尾处为准）
    for di, (speaker, norm_text, key, raw) in enumerate(doc_lines):
        if di in used_doc or len(norm_text) < 4:
            continue
        for i, row in enumerate(rows):
            if i in exact_rows or row["npc"] != speaker or not row["norm"]:
                continue
            short, long_ = sorted((norm_text, row["norm"]), key=len)
            if len(short) >= 4 and short in long_:
                assigned[i] = (key, raw)
                used_doc.add(di)

    # 3) 剩余用相似度兜底
    for di, (speaker, norm_text, key, raw) in enumerate(doc_lines):
        if di in used_doc or len(norm_text) < 6:
            continue
        best_i, best_ratio = None, 0.0
        for i, row in enumerate(rows):
            if i in assigned or row["npc"] != speaker:
                continue
            ratio = difflib.SequenceMatcher(None, norm_text, row["norm"]).ratio()
            if ratio > best_ratio:
                best_i, best_ratio = i, ratio
        if best_i is not None and best_ratio >= 0.72:
            assigned[best_i] = (key, raw)
            used_doc.add(di)

    return assigned, used_doc


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--src", required=True, help="各事件对白文本.docx")
    parser.add_argument("--datas", default="GameConfig/Datas", help="表目录")
    parser.add_argument("--apply", action="store_true", help="写回 xlsx")
    parser.add_argument("--report", default="", help="报告输出路径")
    args = parser.parse_args()

    doc_lines = read_doc(args.src)
    print("文档台词：%d 句，其中带表情 %d 句" %
          (len(doc_lines), sum(1 for d in doc_lines if d[2])))

    xlsx = os.path.join(args.datas, TABLE_NAME)
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.active
    rows = load_rows(ws)
    print("表内节点：%d 行" % len(rows))

    assigned, used_doc = align(doc_lines, rows)
    print("已对齐：%d 行（未匹配的文档句 %d 条）" % (len(assigned), len(doc_lines) - len(used_doc)))

    stat = collections.Counter()
    changed = 0
    for i, row in enumerate(rows):
        key = assigned[i][0] if i in assigned else None
        new_val = key if key else None
        stat[str(new_val)] += 1
        if (row["anim"] or None) != new_val:
            changed += 1
        if args.apply:
            ws.cell(row["row"], COL_ANIM).value = new_val

    print("表情分布：", dict(stat))
    print("实际改动单元格：%d" % changed)

    if args.apply:
        wb.save(xlsx)
        print("已写回：%s" % xlsx)

    if args.report:
        lines = ["# 剧情表情回填报告", "",
                 "- 文档台词：%d 句" % len(doc_lines),
                 "- 表内节点：%d 行" % len(rows),
                 "- 已对齐：%d 行" % len(assigned),
                 "- 表情分布：%s" % dict(stat), "",
                 "## 未匹配到表行的文档台词", ""]
        for di, (speaker, norm_text, key, raw) in enumerate(doc_lines):
            if di not in used_doc:
                lines.append("- [%s/%s] %s" % (speaker, key or "沿用", raw[:70]))
        lines += ["", "## 未拿到表情的表节点（运行时沿用上一句姿势）", ""]
        for i, row in enumerate(rows):
            if i not in assigned:
                lines.append("- `%s` [%s] %s" % (row["id"], row["npc"], row["text"][:60]))
        with open(args.report, "w", encoding="utf-8") as f:
            f.write("\n".join(lines) + "\n")
        print("报告：%s" % args.report)


if __name__ == "__main__":
    main()
