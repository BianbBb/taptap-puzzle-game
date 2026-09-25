"""Workflow safety and validation regression tests."""

import json
import contextlib
import importlib.util
import io
import openpyxl
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from workflow_lib.core import Context, Run, WorkflowError
from workflow_lib import core
from workflow_lib import checks
from workflow_lib import luban
from workflow_lib.unity import decode_response, test_listing, validate_test_result
from workflow import runtime_status


class WorkflowCoreTests(unittest.TestCase):
    def project(self):
        root = Path(self.tmp.name)
        project = root / "GameUnity"
        for folder in ("Assets", "Packages", "ProjectSettings"):
            (project / folder).mkdir(parents=True)
        return project

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def luban_helper(self, data_dir, cache_dir):
        path = Path(__file__).resolve().parents[1] / "skills/luban-dev/scripts/luban_helper.py"
        spec = importlib.util.spec_from_file_location("workflow_test_luban_helper", path)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module.LubanConfigHelper(str(data_dir), str(cache_dir))

    def test_named_sheet_schema_is_detected_and_registered(self):
        data = Path(self.tmp.name) / "Datas"
        (data / "Defines").mkdir(parents=True)
        registry = openpyxl.Workbook()
        sheet = registry.active
        sheet.append(["##var", "full_name", "value_type", "read_schema_from_file", "input"])
        sheet.append(["##", "名称", "类型", "是否读取", "输入"])
        registry.save(data / "Defines/__tables__.xlsx")
        registry.close()

        workbook = openpyxl.Workbook()
        workbook.active.title = "Default"
        named = workbook.create_sheet("Named")
        named.append(["##var", "Id"])
        named.append(["##type", "int"])
        named.append(["##", "ID"])
        workbook.save(data / "named.xlsx")
        workbook.close()

        helper = self.luban_helper(data, Path(self.tmp.name) / "cache")
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertTrue(helper.add_table(
                "TbAudit", [{"name": "Id", "type": "int", "group": "c"}],
                input_file="named.xlsx", sheet_name="Named"))
        workbook = openpyxl.load_workbook(data / "Defines/__tables__.xlsx", data_only=True)
        self.assertTrue(workbook.active.cell(workbook.active.max_row, 4).value)
        workbook.close()

    def test_update_field_controls_multiline_type_rows(self):
        data = Path(self.tmp.name) / "Datas"
        data.mkdir()
        workbook_path = data / "multi.xlsx"
        helper = self.luban_helper(data, Path(self.tmp.name) / "cache")
        helper._get_table_excel_path = lambda *_args: (workbook_path, None)

        def write_sheet():
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["##var", "StepParam", "Other"])
            sheet.append(["##type", "GuideStepParam?", "int"])
            sheet.append(["##type", "int", None])
            sheet.append(["##", "参数", "其他"])
            workbook.save(workbook_path)
            workbook.close()

        write_sheet()
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertTrue(helper.update_field("TbAudit", "StepParam", new_type="int"))
        workbook = openpyxl.load_workbook(workbook_path, data_only=True)
        self.assertEqual([workbook.active.cell(2, 2).value, workbook.active.cell(3, 2).value], ["int", None])
        workbook.close()

        write_sheet()
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertTrue(helper.update_field("TbAudit", "StepParam", new_type="int|string"))
        workbook = openpyxl.load_workbook(workbook_path, data_only=True)
        self.assertEqual([workbook.active.cell(2, 2).value, workbook.active.cell(3, 2).value], ["int", "string"])
        workbook.close()

        write_sheet()
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertFalse(helper.update_field("TbAudit", "StepParam", new_type="int|string|long"))
        workbook = openpyxl.load_workbook(workbook_path, data_only=True)
        self.assertEqual(workbook.active.cell(2, 2).value, "GuideStepParam?")
        workbook.close()

    def test_manifest_mismatch_is_rejected(self):
        with self.assertRaises(WorkflowError):
            luban.validate_manifest({"output/a.bytes": "new"}, {"output/a.bytes": "old"})

    def test_wrong_project_is_blocked(self):
        with self.assertRaises(WorkflowError) as raised:
            Context(Path(self.tmp.name) / "Other")
        self.assertEqual(raised.exception.status, "blocked")

    def test_local_unity_cli_takes_precedence(self):
        project = self.project()
        local = project / "Tools" / "unity.exe"
        local.parent.mkdir()
        local.write_bytes(b"cli")
        with patch.object(core.shutil, "which", return_value=str(Path(self.tmp.name) / "global-unity.exe")):
            context = Context(project)
        self.assertEqual(context.cli, local.resolve())

    def test_unity_cli_falls_back_to_path(self):
        project = self.project()
        global_cli = Path(self.tmp.name) / "global-unity.exe"
        global_cli.write_bytes(b"cli")
        with patch.object(core.shutil, "which", return_value=str(global_cli)):
            context = Context(project)
        self.assertEqual(context.cli, global_cli.resolve())

    def test_new_source_missing_from_solution_is_blocked(self):
        project = self.project()
        (project / "Assets/NewSource.cs").write_text("class NewSource {}", encoding="utf-8")
        (project / "Game.csproj").write_text("<Project></Project>", encoding="utf-8")
        (project / "GameUnity.sln").write_text(
            'Project("{FAKE}") = "Game", "Game.csproj", "{GAME}"\nEndProject\n', encoding="utf-8")
        with self.assertRaises(WorkflowError) as raised:
            checks.solution_inputs(Context(project))
        self.assertEqual(raised.exception.status, "blocked")

    def test_zero_tests_are_blocked(self):
        with self.assertRaises(WorkflowError) as raised:
            test_listing({"tests": []}, "editor", "DGame", "assembly")
        self.assertEqual(raised.exception.status, "blocked")

    def test_stale_test_result_is_rejected(self):
        value = {"status": "completed", "summary": {"total": 1, "passed": 1, "failed": 0,
                 "skipped": 0, "inconclusive": 0}, "results": [{"fullName": "Old.Test"}]}
        with self.assertRaises(WorkflowError):
            validate_test_result(value, {"New.Test"})

    def test_response_parser_rejects_non_json(self):
        with self.assertRaises(WorkflowError):
            decode_response("not json")

    def test_timeout_is_blocked(self):
        run = Run(Context(self.project()), "timeout-test")
        with self.assertRaises(WorkflowError) as raised:
            run.execute([sys.executable, "-c", "import time; time.sleep(0.2)"], timeout=0.01)
        self.assertEqual(raised.exception.status, "blocked")

    def test_luban_isolation_requires_project_tool(self):
        run = Run(Context(self.project()), "luban-test")
        with self.assertRaises(WorkflowError) as raised:
            luban.isolated_export(run)
        self.assertEqual(raised.exception.status, "blocked")

    def test_code_profile_runs_structure_tests_and_build(self):
        context = Context(self.project())
        calls = []

        class FakeRun:
            def __init__(self):
                self.context = context
                self.steps = []

            def step(self, name, function, required=True):
                calls.append(name)
                self.steps.append(name)
                function()
                return {"name": name, "status": "passed", "required": required}

        run = FakeRun()
        original = checks.structure, checks.unit_tests, checks.build
        try:
            checks.structure = lambda context: {"ok": True}
            checks.unit_tests = lambda run: {"tests": 1}
            checks.build = lambda run: {"built": True}
            checks.verify(run, "code", "DGame", "assembly", 1)
        finally:
            checks.structure, checks.unit_tests, checks.build = original
        self.assertEqual(calls, ["workflow-structure", "workflow-unit-tests", "solution-build"])

    def test_luban_manifest_requires_bridge_files_and_outputs(self):
        root = Path(self.tmp.name)
        bridge = root / "GameUnity/Assets/Scripts/HotFix/GameProto"
        for directory in (bridge, root / "GameUnity/Assets/BundleAssets/Configs/Binary",
                          root / "GameUnity/Configs/Json", bridge / "LubanConfig"):
            directory.mkdir(parents=True)
        (bridge / "ConfigSystem.cs").write_text("class ConfigSystem {}", encoding="utf-8")
        with self.assertRaises(WorkflowError) as raised:
            luban._manifest(root)
        self.assertIn("ExternalTypeUtil.cs", str(raised.exception))

    def test_runtime_status_uses_runtime_descriptor(self):
        player = Path(self.tmp.name) / "Player"
        player.mkdir()
        (player / ".unity-pipeline-runtime-port").write_text(
            json.dumps({"workingDirectory": str(Path(self.tmp.name) / "Launcher"),
                        "port": 7900, "evalToken": "token"}), encoding="utf-8")

        class RuntimeClient:
            def raw(self, args):
                self.args = args
                return json.dumps({"IsPlaying": True, "UnityVersion": "2022.3"})

        client = RuntimeClient()
        result = runtime_status(client, player)
        self.assertEqual(Path(result["workingDirectory"]).resolve(), (Path(self.tmp.name) / "Launcher").resolve())
        self.assertTrue(result["status"]["IsPlaying"])

    def test_runtime_status_allows_different_process_working_directory(self):
        player = Path(self.tmp.name) / "Player"
        player.mkdir()
        (player / ".unity-pipeline-runtime-port").write_text(
            json.dumps({"workingDirectory": str(Path(self.tmp.name) / "Other"),
                        "port": 7900, "evalToken": "token"}), encoding="utf-8")
        class RuntimeClient:
            def raw(self, args):
                return json.dumps({"IsPlaying": True})
        result = runtime_status(RuntimeClient(), player)
        self.assertEqual(result["port"], 7900)

    def test_runtime_status_rejects_incomplete_descriptor(self):
        player = Path(self.tmp.name) / "Player"
        player.mkdir()
        (player / ".unity-pipeline-runtime-port").write_text(
            json.dumps({"workingDirectory": str(player), "port": 7900}), encoding="utf-8")
        with self.assertRaises(WorkflowError) as raised:
            runtime_status(object(), player)
        self.assertEqual(raised.exception.status, "blocked")


if __name__ == "__main__":
    unittest.main()
